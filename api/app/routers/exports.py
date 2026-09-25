from fastapi import APIRouter, Depends, HTTPException, Query, Request
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from typing import Optional
import io

from app.database import get_db
from app.models import activities as activity_models
from app.utils.jwt_utils import decode_access_token

router = APIRouter(prefix="", tags=["数据导出"])


def get_current_user_id(request) -> Optional[int]:
    auth_header = request.headers.get("Authorization", "")
    if auth_header.startswith("Bearer "):
        payload = decode_access_token(auth_header[7:])
        if payload and "user_id" in payload:
            return payload["user_id"]
    return None


def require_admin(user_id: int, db: Session):
    from app.models.users import User
    user = db.query(User).filter(User.id == user_id).first()
    if not user or user.role.value != "admin":
        raise HTTPException(status_code=403, detail="无权限")
    return user


@router.get("/export/registrations")
async def export_registrations(
    db: Session = Depends(get_db),
    activity_id: Optional[int] = Query(default=None),
    request: Request = None,
):
    """导出活动报名数据为 Excel"""
    user_id = get_current_user_id(request)
    if not user_id:
        raise HTTPException(status_code=401, detail="未登录")
    require_admin(user_id, db)

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill
    except ImportError:
        return {"error": "请安装 openpyxl: pip install openpyxl"}

    q = db.query(activity_models.ActivityRegistration)
    if activity_id:
        q = q.filter(activity_models.ActivityRegistration.activity_id == activity_id)

    registrations = q.order_by(activity_models.ActivityRegistration.registered_at.desc()).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "报名数据"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center")

    headers = ["活动名称", "报名类型", "姓名", "联系电话", "所在区域", "意向参与内容", "备注", "报名时间"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    activity_map = {}
    for r in registrations:
        act = activity_map.get(r.activity_id)
        if not act:
            act = db.query(activity_models.Activity).filter(activity_models.Activity.id == r.activity_id).first()
            activity_map[r.activity_id] = act

        row_num = ws.max_row + 1
        ws.cell(row=row_num, column=1, value=act.title if act else "")
        ws.cell(row=row_num, column=2, value="个人" if r.register_type == "personal" else "商户")
        ws.cell(row=row_num, column=3, value=r.name)
        ws.cell(row=row_num, column=4, value=r.phone)
        ws.cell(row=row_num, column=5, value=r.area)
        ws.cell(row=row_num, column=6, value=r.interest or "")
        ws.cell(row=row_num, column=7, value=r.remark or "")
        ws.cell(row=row_num, column=8, value=r.registered_at.strftime("%Y-%m-%d %H:%M") if r.registered_at else "")

    for col_cells in ws.columns:
        max_len = 0
        col_letter = col_cells[0].column_letter
        for cell in col_cells:
            try:
                if cell.value and len(str(cell.value)) > max_len:
                    max_len = len(str(cell.value))
            except Exception:
                pass
        if max_len > 0:
            ws.column_dimensions[col_letter].width = min(max_len + 2, 40)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"活动报名数据_{activity_id or '全部'}.xlsx"
    import urllib.parse
    safe_filename = urllib.parse.quote(filename)
    return StreamingResponse(
        iter([buf.read()]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{safe_filename}"},
    )
